from selenium import webdriver
from selenium.webdriver.support import expected_conditions as EC
from selenium.webdriver.support.ui import WebDriverWait
from selenium.common.exceptions import TimeoutException
from selenium.webdriver.common.by import By

from openpyxl import load_workbook
import openpyxl
import time

# os.chdir('C:\Users\rezaphonna\Documents\belajar\office')
# wb = openpyxl.load_workbook('data.xlsx')

wb = load_workbook(filename="C:\\Users\\rezaphonna\\Documents\\belajar\\office\\data.xlsx")

sheetRange = wb['Sheet1']

options = webdriver.ChromeOptions()
options.add_experimental_option('detach',True)
driver = webdriver.Chrome(options=options)
driver.maximize_window()
driver.implicitly_wait(10)
driver.get("https://demoqa.com/webtables")

"""
driver = webdriver.Chrome()
driver.get("https://demoqa.com/webtables")
driver.maximize_window()
driver.implicitly_wait(10)
"""
#Looping input

i = 2
while i <=len(sheetRange['A']):
    #BAca baris file excel
    Firstname = sheetRange['A'+str(i)].value
    Lastname = sheetRange['B'+str(i)].value
    Email = sheetRange['C'+str(i)].value
    Age = sheetRange['D'+str(i)].value
    Salary = sheetRange['E'+str(i)].value
    Departemen = sheetRange['F'+str(i)].value

    driver.find_element(By.ID,"addNewRecordButton").click()

    try:
        element = WebDriverWait(driver,10).until(EC.visibility_of_element_located((By.XPATH, "/html/body/div[5]/div/div")))

        driver.find_element(By.ID,"firstName").send_keys(Firstname)
        driver.find_element(By.ID,"lastName").send_keys(Lastname) 
        driver.find_element(By.ID,"userEmail").send_keys(Email) 
        driver.find_element(By.ID,"age").send_keys(Age) 
        driver.find_element(By.ID,"salary").send_keys(Salary)
        driver.find_element(By.ID,"department").send_keys(Departemen)
        driver.find_element(By.ID,"submit").click()

        print("elemen ditemukan")

    except TimeoutException:
        print("form ga muncul")
        pass

    time.sleep(1)
    i = i + 1

print("Beres")
